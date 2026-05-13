import csv
import io
import re

from flask import Blueprint, jsonify, request, send_file
from openpyxl import Workbook, load_workbook

from routes.admin import admin_required
from models.university import University
from models.college import College
from models.industry import Industry
from models.project_title import ProjectTitle
from models.major import Major
from models.degree import Degree
from models.department import Department
from models.work_keyword import WorkKeyword
from models.career_objective import CareerObjective
from models.resume_keyword import ResumeKeyword
from models.audit_log import AuditLog

bulk_import_bp = Blueprint('bulk_import', __name__)


def _log_audit(admin_user, action, target_type='', target_id='', details=None):
    try:
        from flask import request as _req
        AuditLog(
            adminUser=admin_user.id,
            action=action,
            targetType=target_type,
            targetId=str(target_id) if target_id else '',
            details=details or {},
            ipAddress=_req.remote_addr or '',
        ).save()
    except Exception:
        pass

ENTITY_CONFIG = {
    'universities': {
        'filename': 'universities_template.xlsx',
        'sheet': 'Universities',
        'headers': ['University Name', 'City', 'District', 'State', 'Website', 'Logo Path'],
        'sample': ['Example University', 'Surat', 'Surat', 'Gujarat', 'https://example.edu', ''],
    },
    'colleges': {
        'filename': 'colleges_template.xlsx',
        'sheet': 'Colleges',
        'headers': [
            'College Name',
            'College City',
            'College Tehsil',
            'College District',
            'College State',
            'College Website',
            'College Logo Path',
        ],
        'sample': [
            'Example College',
            'Surat',
            'Choryasi',
            'Surat',
            'Gujarat',
            'https://example.edu/college',
            '',
        ],
    },
    'industries': {
        'filename': 'industries_template.xlsx',
        'sheet': 'Industries',
        'headers': ['Industry Name', 'City', 'District', 'State', 'GST Number', 'Website', 'Type', 'Sub-Type', 'Details'],
        'sample': ['Example Industry', 'Surat', 'Surat', 'Gujarat', '24AAAAA0000A1Z5', 'https://example.com', 'Services', 'Software Development', ''],
    },
    'project-titles': {
        'filename': 'project_titles_template.xlsx',
        'sheet': 'Project Titles',
        'headers': ['Title', 'Major Name', 'Degree Name'],
        'sample': ['Sample Project Title', 'Computer Science', 'BCA'],
    },
    'degree-structure': {
        'filename': 'degree_structure_template.xlsx',
        'sheet': 'Degree Structure',
        'headers': ['Degree', 'Department', 'Major 1', 'Major 2', 'Major 3'],
        'sample': ['Bachelor of Technology', 'Department of Computer Engineering', 'Computer Engineering', 'Information Technology', ''],
    },
    'work-keywords': {
        'filename': 'work_keywords_template.xlsx',
        'sheet': 'Work Keywords',
        'headers': ['Keyword', 'Industry Type', 'Job Profile', 'Sort Order', 'Major Name'],
        'sample': ['Sale of dairy products', 'Sales', 'Sales Executive', '10', 'Commerce'],
    },
    'career-objectives': {
        'filename': 'career_objectives_template.xlsx',
        'sheet': 'Career Objectives',
        'headers': ['Text', 'Major Name', 'Active'],
        'sample': ['To become a Python developer.', 'Computer Science', 'Yes'],
    },
    'resume-keywords': {
        'filename': 'resume_keywords_template.xlsx',
        'sheet': 'Resume Keywords',
        'headers': ['Keyword', 'Category', 'Major Name', 'Industry Type', 'Job Profile', 'Sort Order', 'Active'],
        'sample': ['Developed REST APIs', 'experience', 'Computer Science', 'Software', 'Backend Developer', '10', 'Yes'],
    },
}

HEADER_ALIASES = {
    'universities': {
        'university name': 'name',
        'name': 'name',
        'city': 'villageCityName',
        'village city name': 'villageCityName',
        'district': 'district',
        'state': 'state',
        'website': 'website',
        'logo path': 'logo',
        'logo': 'logo',
    },
    'colleges': {
        'college name': 'name',
        'name': 'name',
        'city': 'villageCityName',
        'village city name': 'villageCityName',
        'college city': 'villageCityName',
        'college tehsil': 'tehsil',
        'tehsil': 'tehsil',
        'district': 'district',
        'college district': 'district',
        'state': 'state',
        'college state': 'state',
        'website': 'website',
        'college website': 'website',
        'logo path': 'logo',
        'logo': 'logo',
        'college logo path': 'logo',
        'college logo': 'logo',
    },
    'industries': {
        'industry name': 'name',
        'name': 'name',
        'city': 'villageCityName',
        'village city name': 'villageCityName',
        'district': 'district',
        'state': 'state',
        'gst number': 'gstNumber',
        'website': 'website',
        'type': 'industryType',
        'sub type': 'industrySubType',
        'sub-type': 'industrySubType',
        'details': 'industryDetails',
    },
    'project-titles': {
        'title': 'title',
        'major name': 'major',
        'major': 'major',
        'degree name': 'degree',
        'degree': 'degree',
    },
    'degree-structure': {
        'degree': 'degree',
        'department': 'department',
    },
    'work-keywords': {
        'keyword': 'keyword',
        'industry type': 'industryType',
        'job profile': 'jobProfile',
        'sort order': 'sortOrder',
        'major name': 'major',
        'major': 'major',
    },
    'career-objectives': {
        'text': 'text',
        'major name': 'major',
        'major': 'major',
        'active': 'isActive',
        'is active': 'isActive',
    },
    'resume-keywords': {
        'keyword': 'keyword',
        'category': 'category',
        'major name': 'major',
        'major': 'major',
        'industry type': 'industryType',
        'job profile': 'jobProfile',
        'sort order': 'sortOrder',
        'active': 'isActive',
        'is active': 'isActive',
    },
}


def _normalize_entity_key(entity):
    key = str(entity or '').strip().lower().replace('_', '-')
    if key == 'project-titles' or key == 'project-title':
        return 'project-titles'
    if key in ('degree-structure', 'degreestructure', 'degree_structure', 'majors'):
        return 'degree-structure'
    if key in ('work-keyword', 'workkeywords', 'keywords', 'work-keywords'):
        return 'work-keywords'
    return key


def _normalize_header(value):
    return re.sub(r'[^a-z0-9]+', ' ', str(value or '').strip().lower()).strip()


def _normalize_text(value):
    text = re.sub(r'[^\w\s]+', ' ', str(value or ''))
    return re.sub(r'\s+', ' ', text).strip().casefold()


def _obj_id(doc):
    if not doc:
        return None
    direct_id = getattr(doc, 'id', None)
    if direct_id is not None:
        return str(direct_id)
    raw_data = getattr(doc, '_data', None) or {}
    raw_id = raw_data.get('id')
    return str(raw_id) if raw_id is not None else None


def _find_by_normalized_name(model_cls, name):
    target = _normalize_text(name)
    if not target:
        return None
    for doc in model_cls.objects():
        if _normalize_text(getattr(doc, 'name', '')) == target:
            return doc
    return None


def _import_degree_structure(rows, raw_rows):
    """
    Import Degree -> Department -> Major(s) from client rows.
    raw_rows is used to read dynamic Major columns.
    """
    imported_degrees = 0
    imported_departments = 0
    imported_majors = 0
    skipped = 0
    errors = []
    degree_cache = {}
    department_cache = {}
    last_degree_name = ''
    last_department_name = ''

    for index, (row, raw_row) in enumerate(zip(rows, raw_rows), start=2):
        degree_name = _resolve_required_text(row, 'degree') or last_degree_name
        if not degree_name:
            errors.append({'row': index, 'reason': 'Degree is required'})
            continue
        last_degree_name = degree_name

        degree_target = _normalize_text(degree_name)
        degree = degree_cache.get(degree_target)
        if not degree:
            degree = _find_by_normalized_name(Degree, degree_name)
            if not degree:
                degree = Degree(name=degree_name, isActive=True)
                degree.save()
                imported_degrees += 1
            degree_cache[degree_target] = degree

        department_name = _resolve_required_text(row, 'department') or last_department_name
        department = None
        if department_name:
            last_department_name = department_name
            dept_target = _normalize_text(department_name)
            cache_key = (str(_obj_id(degree) or ''), dept_target)
            department = department_cache.get(cache_key)
            if not department:
                department = next(
                    (
                        doc for doc in Department.objects()
                        if _normalize_text(getattr(doc, 'name', '')) == dept_target
                        and str(_obj_id(getattr(doc, 'degree', None)) or '') == str(_obj_id(degree) or '')
                    ),
                    None,
                )
            if not department:
                department = Department(
                    name=department_name,
                    degree=degree,
                    approvalStatus='approved',
                    isActive=True,
                )
                department.save()
                imported_departments += 1
            department_cache[cache_key] = department

        major_names = []
        for key in (raw_row or {}):
            normalized_key = _normalize_header(key)
            if normalized_key == 'major' or (normalized_key.startswith('major') and normalized_key != 'major name'):
                value = str(raw_row[key] or '').strip()
                if value:
                    major_names.append(value)

        if not major_names:
            errors.append({'row': index, 'reason': f'At least one major is required (Degree: {degree_name})'})
            continue

        for major_name in major_names:
            major_target = _normalize_text(major_name)
            existing_major = next(
                (
                    doc for doc in Major.objects()
                    if _normalize_text(getattr(doc, 'name', '')) == major_target
                    and str(_obj_id(getattr(doc, 'degree', None)) or '') == str(_obj_id(degree) or '')
                ),
                None,
            )
            if existing_major:
                skipped += 1
                continue

            Major(
                name=major_name,
                degree=degree,
                department=department,
                isActive=True,
            ).save()
            imported_majors += 1

    return imported_degrees, imported_departments, imported_majors, skipped, errors
def _build_row_maps(rows, entity_key):
    aliases = HEADER_ALIASES[entity_key]
    mapped_rows = []
    for row in rows:
        mapped = {}
        for key, value in row.items():
            normalized = aliases.get(_normalize_header(key))
            if normalized:
                mapped[normalized] = value
        mapped_rows.append(mapped)
    return mapped_rows


def _load_rows_from_upload(file_storage):
    filename = (file_storage.filename or '').lower()
    payload = file_storage.read()
    if filename.endswith('.csv'):
        text = payload.decode('utf-8-sig')
        reader = csv.DictReader(io.StringIO(text))
        headers = reader.fieldnames or []
        rows = list(reader)
        return headers, rows

    workbook = load_workbook(io.BytesIO(payload), data_only=True)
    worksheet = workbook.active
    rows = list(worksheet.iter_rows(values_only=True))
    if not rows:
        return [], []
    headers = [str(value or '').strip() for value in rows[0]]
    data_rows = []
    for values in rows[1:]:
        data_rows.append({headers[index]: values[index] if index < len(values) else None for index in range(len(headers))})
    return headers, data_rows


def _resolve_required_text(row, key):
    return str(row.get(key, '') or '').strip()


def _resolve_optional_bool(row, key, default=True):
    value = row.get(key, None)
    if value is None:
        return default
    if isinstance(value, bool):
        return value
    text = str(value).strip().lower()
    if not text:
        return default
    if text in ('1', 'true', 'yes', 'y', 'active', 'on'):
        return True
    if text in ('0', 'false', 'no', 'n', 'inactive', 'off'):
        return False
    return default


def _import_universities(rows):
    imported = 0
    skipped = 0
    errors = []
    for index, row in enumerate(rows, start=2):
        name = _resolve_required_text(row, 'name')
        if not name:
            errors.append({'row': index, 'reason': 'University Name is required'})
            continue

        existing = _find_by_normalized_name(University, name)
        if existing:
            skipped += 1
            continue

        University(
            name=name,
            villageCityName=_resolve_required_text(row, 'villageCityName'),
            district=_resolve_required_text(row, 'district'),
            state=_resolve_required_text(row, 'state'),
            website=_resolve_required_text(row, 'website'),
            logo=_resolve_required_text(row, 'logo'),
            isVerified=True,
            isActive=True,
            approvalStatus='approved',
        ).save()
        imported += 1
    return imported, skipped, errors


def _import_colleges(rows, university):
    """
    Import colleges and link them all to the given university object.
    university is already resolved before calling this function.
    """
    imported = 0
    skipped = 0
    errors = []
    for index, row in enumerate(rows, start=2):
        name = _resolve_required_text(row, 'name')
        if not name:
            errors.append({'row': index, 'reason': 'College Name is required'})
            continue

        # Check if college with same name already exists under this university
        existing = next(
            (
                doc for doc in College.objects()
                if _normalize_text(getattr(doc, 'name', '')) == _normalize_text(name)
                and str(_obj_id(getattr(doc, 'university', None))) == _obj_id(university)
            ),
            None,
        )
        if existing:
            skipped += 1
            continue

        College(
            name=name,
            university=university,
            villageCityName=_resolve_required_text(row, 'villageCityName'),
            tehsil=_resolve_required_text(row, 'tehsil'),
            district=_resolve_required_text(row, 'district'),
            state=_resolve_required_text(row, 'state'),
            website=_resolve_required_text(row, 'website'),
            logo=_resolve_required_text(row, 'logo'),
            isVerified=True,
            isActive=True,
            approvalStatus='approved',
        ).save()
        imported += 1
    return imported, skipped, errors


def _import_industries(rows):
    imported = 0
    skipped = 0
    errors = []
    for index, row in enumerate(rows, start=2):
        name = _resolve_required_text(row, 'name')
        if not name:
            errors.append({'row': index, 'reason': 'Industry Name is required'})
            continue

        existing = _find_by_normalized_name(Industry, name)
        if existing:
            skipped += 1
            continue

        Industry(
            name=name,
            villageCityName=_resolve_required_text(row, 'villageCityName'),
            district=_resolve_required_text(row, 'district'),
            state=_resolve_required_text(row, 'state'),
            gstNumber=_resolve_required_text(row, 'gstNumber'),
            website=_resolve_required_text(row, 'website'),
            industryType=_resolve_required_text(row, 'industryType'),
            industrySubType=_resolve_required_text(row, 'industrySubType'),
            industryDetails=_resolve_required_text(row, 'industryDetails'),
            isVerified=True,
            isActive=True,
            approvalStatus='approved',
        ).save()
        imported += 1
    return imported, skipped, errors


def _find_major(name, degree_name=''):
    target_name = _normalize_text(name)
    target_degree = _normalize_text(degree_name)
    for major in Major.objects():
        if _normalize_text(getattr(major, 'name', '')) != target_name:
            continue
        if target_degree:
            degree = getattr(major, 'degree', None)
            degree_name_value = getattr(degree, 'name', '') if degree else ''
            if _normalize_text(degree_name_value) != target_degree:
                continue
        return major
    return None


def _import_project_titles(rows):
    imported = 0
    skipped = 0
    errors = []
    for index, row in enumerate(rows, start=2):
        title = _resolve_required_text(row, 'title')
        major_name = _resolve_required_text(row, 'major')
        degree_name = _resolve_required_text(row, 'degree')
        if not title or not major_name:
            errors.append({'row': index, 'reason': 'Title and Major Name are required'})
            continue

        major = _find_major(major_name, degree_name)
        if not major:
            errors.append({'row': index, 'reason': f'Unknown major: {major_name}'})
            continue

        degree = getattr(major, 'degree', None)
        if degree_name:
            resolved_degree = _find_by_normalized_name(Degree, degree_name)
            if resolved_degree:
                degree = resolved_degree

        existing = next(
            (
                doc for doc in ProjectTitle.objects()
                if _normalize_text(getattr(doc, 'title', '')) == _normalize_text(title)
                and str(_obj_id(getattr(doc, 'major', None))) == _obj_id(major)
            ),
            None,
        )
        if existing:
            skipped += 1
            continue

        ProjectTitle(
            title=title,
            major=major,
            degree=degree,
            isActive=True,
        ).save()
        imported += 1
    return imported, skipped, errors


def _import_work_keywords(rows):
    imported = 0
    skipped = 0
    errors = []
    for index, row in enumerate(rows, start=2):
        keyword = _resolve_required_text(row, 'keyword')
        if not keyword:
            errors.append({'row': index, 'reason': 'Keyword is required'})
            continue

        industry_type = _resolve_required_text(row, 'industryType')
        job_profile = _resolve_required_text(row, 'jobProfile')
        sort_order_raw = _resolve_required_text(row, 'sortOrder')
        try:
            sort_order = int(sort_order_raw) if sort_order_raw else 0
        except ValueError:
            sort_order = 0

        # Resolve major (optional)
        major_name = _resolve_required_text(row, 'major')
        major_doc = None
        if major_name:
            major_doc = _find_major(major_name)
            if not major_doc:
                errors.append({'row': index, 'reason': f'Major "{major_name}" not found'})
                continue

        existing = next(
            (
                doc for doc in WorkKeyword.objects()
                if _normalize_text(getattr(doc, 'keyword', '')) == _normalize_text(keyword)
                and _normalize_text(getattr(doc, 'industryType', '')) == _normalize_text(industry_type)
                and _normalize_text(getattr(doc, 'jobProfile', '')) == _normalize_text(job_profile)
                and _obj_id(getattr(doc, 'major', None)) == _obj_id(major_doc)  # Also compare major
            ),
            None,
        )
        if existing:
            skipped += 1
            continue

        WorkKeyword(
            keyword=keyword,
            industryType=industry_type,
            jobProfile=job_profile,
            major=major_doc,
            sortOrder=sort_order,
            isActive=True,
        ).save()
        imported += 1
    return imported, skipped, errors


def _import_career_objectives(rows):
    imported = 0
    skipped = 0
    errors = []
    for index, row in enumerate(rows, start=2):
        text = _resolve_required_text(row, 'text')
        if not text:
            errors.append({'row': index, 'reason': 'Text is required'})
            continue

        major_name = _resolve_required_text(row, 'major')
        major_doc = None
        if major_name:
            major_doc = _find_major(major_name)
            if not major_doc:
                errors.append({'row': index, 'reason': f'Major "{major_name}" not found'})
                continue

        existing = next(
            (
                doc for doc in CareerObjective.objects()
                if _normalize_text(getattr(doc, 'text', '')) == _normalize_text(text)
                and _obj_id(getattr(doc, 'major', None)) == _obj_id(major_doc)
            ),
            None,
        )
        if existing:
            skipped += 1
            continue

        CareerObjective(
            text=text,
            major=major_doc,
            isActive=_resolve_optional_bool(row, 'isActive', True),
            approvalStatus='approved',
            rejectionReason='',
        ).save()
        imported += 1
    return imported, skipped, errors


def _import_resume_keywords(rows):
    imported = 0
    skipped = 0
    errors = []
    valid_categories = {'experience', 'project', 'volunteering', 'skill', 'technical'}
    for index, row in enumerate(rows, start=2):
        keyword = _resolve_required_text(row, 'keyword')
        category = _resolve_required_text(row, 'category').lower()
        if not keyword or not category:
            errors.append({'row': index, 'reason': 'Keyword and Category are required'})
            continue
        if category not in valid_categories:
            errors.append({'row': index, 'reason': 'Category must be one of experience, project, volunteering, skill, technical'})
            continue

        major_name = _resolve_required_text(row, 'major')
        major_doc = None
        if major_name:
            major_doc = _find_major(major_name)
            if not major_doc:
                errors.append({'row': index, 'reason': f'Major "{major_name}" not found'})
                continue

        sort_order_raw = _resolve_required_text(row, 'sortOrder')
        try:
            sort_order = int(sort_order_raw) if sort_order_raw else 0
        except ValueError:
            sort_order = 0

        industry_type = _resolve_required_text(row, 'industryType')
        job_profile = _resolve_required_text(row, 'jobProfile')

        existing = next(
            (
                doc for doc in ResumeKeyword.objects()
                if _normalize_text(getattr(doc, 'keyword', '')) == _normalize_text(keyword)
                and _normalize_text(getattr(doc, 'category', '')) == _normalize_text(category)
                and _obj_id(getattr(doc, 'major', None)) == _obj_id(major_doc)
                and _normalize_text(getattr(doc, 'industryType', '')) == _normalize_text(industry_type)
                and _normalize_text(getattr(doc, 'jobProfile', '')) == _normalize_text(job_profile)
            ),
            None,
        )
        if existing:
            skipped += 1
            continue

        ResumeKeyword(
            keyword=keyword,
            category=category,
            major=major_doc,
            industryType=industry_type,
            jobProfile=job_profile,
            sortOrder=sort_order,
            isActive=_resolve_optional_bool(row, 'isActive', True),
        ).save()
        imported += 1
    return imported, skipped, errors


IMPORT_HANDLERS = {
    'universities': _import_universities,
    'colleges': _import_colleges,
    'degree-structure': None,
    'industries': _import_industries,
    'project-titles': _import_project_titles,
    'work-keywords': _import_work_keywords,
    'career-objectives': _import_career_objectives,
    'resume-keywords': _import_resume_keywords,
}


@bulk_import_bp.route('/<entity>', methods=['POST'])
@admin_required
def import_entity(current_user, entity):
    entity_key = _normalize_entity_key(entity)
    if entity_key not in IMPORT_HANDLERS:
        return jsonify({'message': 'Unsupported entity'}), 400

    if entity_key == 'degree-structure':
        file_storage = request.files.get('file')
        if not file_storage:
            return jsonify({'message': 'file is required'}), 400

        filename = (file_storage.filename or '').lower()
        if not filename.endswith(('.csv', '.xlsx')):
            return jsonify({'message': 'Only CSV and XLSX files are supported'}), 400

        try:
            headers, raw_rows = _load_rows_from_upload(file_storage)
            if not headers or not raw_rows:
                return jsonify({'message': 'Uploaded file is empty'}), 400

            mapped_rows = _build_row_maps(raw_rows, entity_key)
            imported_degrees, imported_departments, imported_majors, skipped, errors = _import_degree_structure(mapped_rows, raw_rows)

            _log_audit(current_user, 'bulk_import', 'degree-structure', '',
                       {
                           'imported_degrees': imported_degrees,
                           'imported_departments': imported_departments,
                           'imported_majors': imported_majors,
                           'skipped': skipped,
                       })

            return jsonify({
                'entity': entity_key,
                'imported_degrees': imported_degrees,
                'imported_departments': imported_departments,
                'imported_majors': imported_majors,
                'imported': imported_majors,
                'skipped': skipped,
                'errors': errors,
            })
        except Exception as exc:
            return jsonify({'message': f'Import failed: {str(exc)}'}), 500

    file_storage = request.files.get('file')
    if not file_storage:
        return jsonify({'message': 'file is required'}), 400

    filename = (file_storage.filename or '').lower()
    if not filename.endswith(('.csv', '.xlsx')):
        return jsonify({'message': 'Only CSV and XLSX files are supported'}), 400

    try:
        headers, raw_rows = _load_rows_from_upload(file_storage)
        if not headers or not raw_rows:
            return jsonify({'message': 'Uploaded file is empty'}), 400

        mapped_rows = _build_row_maps(raw_rows, entity_key)

        # Special-case colleges: admin must supply a university_id and
        # imported colleges will be linked to that university.
        if entity_key == 'colleges':
            university_id = request.form.get('university_id')
            if not university_id:
                return jsonify({'message': 'university_id is required for colleges import'}), 400
            university = University.objects(id=university_id).first()
            if not university:
                return jsonify({'message': 'University not found'}), 400
            imported, skipped, errors = _import_colleges(mapped_rows, university)
            _log_audit(current_user, 'bulk_import', 'colleges', '',
                       {'imported': imported, 'skipped': skipped})
            return jsonify({
                'entity': entity_key,
                'imported': imported,
                'skipped': skipped,
                'errors': errors,
            })

        imported, skipped, errors = IMPORT_HANDLERS[entity_key](mapped_rows)
        if entity_key == 'universities':
            _log_audit(current_user, 'bulk_import', 'universities', '',
                       {'imported': imported, 'skipped': skipped})
        elif entity_key == 'colleges':
            _log_audit(current_user, 'bulk_import', 'colleges', '',
                       {'imported': imported, 'skipped': skipped})
        elif entity_key == 'industries':
            _log_audit(current_user, 'bulk_import', 'industries', '',
                       {'imported': imported, 'skipped': skipped})
        elif entity_key == 'project-titles':
            _log_audit(current_user, 'bulk_import', 'project-titles', '',
                       {'imported': imported, 'skipped': skipped})
        elif entity_key == 'work-keywords':
            _log_audit(current_user, 'bulk_import', 'work-keywords', '',
                       {'imported': imported, 'skipped': skipped})
        elif entity_key == 'career-objectives':
            _log_audit(current_user, 'bulk_import', 'career-objectives', '',
                       {'imported': imported, 'skipped': skipped})
        elif entity_key == 'resume-keywords':
            _log_audit(current_user, 'bulk_import', 'resume-keywords', '',
                       {'imported': imported, 'skipped': skipped})
        return jsonify({
            'entity': entity_key,
            'imported': imported,
            'skipped': skipped,
            'errors': errors,
        })
    except Exception as exc:
        return jsonify({'message': f'Import failed: {str(exc)}'}), 500


@bulk_import_bp.route('/template/<entity>', methods=['GET'])
@admin_required
def download_template(current_user, entity):
    entity_key = _normalize_entity_key(entity)
    config = ENTITY_CONFIG.get(entity_key)
    if not config:
        return jsonify({'message': 'Unsupported entity'}), 400

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = config['sheet']
    worksheet.append(config['headers'])
    worksheet.append(config['sample'])

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)

    return send_file(
        buffer,
        as_attachment=True,
        download_name=config['filename'],
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
