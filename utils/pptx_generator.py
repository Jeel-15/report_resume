"""
utils/pptx_generator.py

Build a presentation-style PPTX from assignment sections.
Uses python-pptx and keeps the output simple, readable, and export-friendly.
"""
import io
import re

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt


FONT_HEADLINE = 'Aptos Display'
FONT_BODY = 'Aptos'
COLOR_NAVY = RGBColor(0x10, 0x1E, 0x3A)
COLOR_INK = RGBColor(0x17, 0x21, 0x37)
COLOR_MUTED = RGBColor(0x5B, 0x6B, 0x8A)
COLOR_SOFT = RGBColor(0xF5, 0xF7, 0xFC)
COLOR_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
COLOR_ACCENT = RGBColor(0x2F, 0x6B, 0xFF)
COLOR_ACCENT_2 = RGBColor(0xFF, 0x7A, 0x59)
COLOR_ACCENT_3 = RGBColor(0x22, 0xC5, 0xA1)


def _clean_text(value):
    text = '' if value is None else str(value)
    text = text.replace('\r\n', '\n').replace('\r', '\n')
    text = re.sub(r'\n{3,}', '\n\n', text)
    return text.strip()


def _strip_slide_prefix(text):
    value = _clean_text(text)
    value = re.sub(r'(?i)^\s*slide\s*\d+\s*:\s*', '', value)
    return value.strip()


def _strip_inline_markdown(text):
    value = _clean_text(text)
    if not value:
        return value

    # Remove markdown heading markers while keeping the text itself.
    value = re.sub(r'^\s{0,3}#{1,6}\s*', '', value)
    value = re.sub(r'\*\*(.+?)\*\*', r'\1', value)
    value = re.sub(r'(?<!\w)\*(.+?)\*(?!\w)', r'\1', value)
    value = re.sub(r'`([^`]+)`', r'\1', value)
    value = re.sub(r'\[(.*?)\]\((.*?)\)', r'\1', value)
    value = re.sub(r'\[(.*?)\]', r'\1', value)
    value = value.replace('—', '-').replace('–', '-')
    value = re.sub(r'[\*_]{2,}', '', value)
    value = re.sub(r'\s+', ' ', value)
    return value.strip()


def _extract_speaker_notes(content_md):
    text = _clean_text(content_md)
    if not text:
        return '', ''

    lines = text.split('\n')
    visible_lines = []
    notes_lines = []
    in_notes = False

    for raw_line in lines:
        line = raw_line.rstrip()
        stripped = line.strip()
        if not stripped:
            if in_notes:
                notes_lines.append('')
            else:
                visible_lines.append('')
            continue

        notes_header = re.match(r'(?i)^\s*(speaker\s*notes?|notes)\s*:\s*(.*)$', stripped)
        if notes_header:
            in_notes = True
            tail = notes_header.group(2).strip()
            if tail:
                notes_lines.append(tail)
            continue

        if in_notes:
            notes_lines.append(line)
        else:
            visible_lines.append(line)

    visible = '\n'.join(visible_lines).strip()
    notes = '\n'.join(notes_lines).strip()
    return visible, notes


def _markdown_to_slide_blocks(markdown_text, slide_title=''):
    """Convert section markdown into structured slide blocks.

    Returns a list of tuples: (kind, text) where kind is one of
    'label', 'bullet', 'number', or 'paragraph'.
    """
    text, _ = _extract_speaker_notes(markdown_text)
    if not text:
        return []

    blocks = []
    slide_title_norm = _strip_slide_prefix(slide_title).lower()

    for raw_line in text.split('\n'):
        line = raw_line.strip()
        if not line:
            continue

        normalized = _strip_inline_markdown(line)
        if not normalized:
            continue

        # Ignore repeated slide title / speaker note scaffolding already handled by the deck layout.
        if re.match(r'(?i)^\s*slide\s*\d+\s*:\s*', normalized):
            candidate = _strip_slide_prefix(normalized)
            if not candidate or candidate.lower() == slide_title_norm:
                continue
            normalized = candidate


        bullet_match = re.match(r'^[-*•]\s+(.+)', line)
        if bullet_match:
            bullet_text = _strip_inline_markdown(bullet_match.group(1))
            if bullet_text:
                blocks.append(('bullet', bullet_text))
            continue

        num_match = re.match(r'^(\d+)\.\s+(.+)', line)
        if num_match:
            number_text = _strip_inline_markdown(num_match.group(2))
            if number_text:
                blocks.append(('number', f'{num_match.group(1)}. {number_text}'))
            continue

        if line.startswith('|'):
            parts = [part.strip() for part in line.strip('|').split('|') if part.strip()]
            if parts:
                blocks.append(('paragraph', '  '.join(_strip_inline_markdown(part) for part in parts)))
            continue

        # Convert bare emphasis labels like "Headline:" / "Bullet Points:" into clean labels.
        if normalized.endswith(':') and len(normalized) <= 80:
            blocks.append(('label', normalized))
            continue

        blocks.append(('paragraph', normalized))

    return blocks


def _chunk_lines(lines, max_lines=7, max_chars=700):
    chunks = []
    current = []
    current_chars = 0
    for line in lines:
        line_len = len(line)
        if current and (len(current) >= max_lines or current_chars + line_len > max_chars):
            chunks.append(current)
            current = []
            current_chars = 0
        current.append(line)
        current_chars += line_len
    if current:
        chunks.append(current)
    return chunks or [[]]


def _add_top_bar(slide, prs, color=RGBColor(0x2A, 0x4C, 0xD6)):
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        0,
        0,
        prs.slide_width,
        Inches(0.25),
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = color
    bar.line.fill.background()


def _set_background(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def _add_footer(slide, label, slide_number, total_slides):
    footer = slide.shapes.add_textbox(Inches(0.7), Inches(6.92), Inches(11.2), Inches(0.24))
    tf = footer.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = label
    run.font.name = FONT_BODY
    run.font.size = Pt(9)
    run.font.color.rgb = COLOR_MUTED

    counter = slide.shapes.add_textbox(Inches(11.85), Inches(6.88), Inches(0.85), Inches(0.24))
    tf = counter.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    run = p.add_run()
    run.text = f'{slide_number}/{total_slides}'
    run.font.name = FONT_BODY
    run.font.size = Pt(9)
    run.font.color.rgb = COLOR_MUTED


def _add_section_tag(slide, text, left=11.15, top=0.48, width=1.5, height=0.34, color=COLOR_ACCENT):
    tag = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    tag.fill.solid()
    tag.fill.fore_color.rgb = color
    tag.line.fill.background()
    tf = tag.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.name = FONT_BODY
    run.font.size = Pt(10)
    run.font.bold = True
    run.font.color.rgb = COLOR_WHITE


def _add_card(slide, left, top, width, height, fill=COLOR_WHITE, line=RGBColor(0xD8, 0xE1, 0xF2), radius=False):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    card = slide.shapes.add_shape(shape_type, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = fill
    card.line.color.rgb = line
    return card


def _add_textbox(slide, left, top, width, height, text, font_size=24, bold=False, color=RGBColor(0x1A, 0x1A, 0x2E), align=PP_ALIGN.LEFT, font_name=FONT_BODY):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def generate_pptx_from_sections(sections, metadata):
    """
    Create a PPTX deck from assignment sections.

    Args:
        sections: list of dicts with section_key, title, content_markdown, sort_order
        metadata: dict with title, assignmentType, studentName, universityName, degreeName, wordCount

    Returns:
        io.BytesIO with the PPTX file.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    title = str(metadata.get('title', 'Presentation') or 'Presentation').strip()
    assignment_type_label = {
        'essay': 'Essay',
        'research_paper': 'Research Paper',
        'lab_report': 'Lab Report',
        'case_study': 'Case Study',
        'presentation_script': 'Presentation',
    }.get(str(metadata.get('assignmentType', '') or ''), 'Presentation')
    student_name = str(metadata.get('studentName', '') or '').strip()
    university_name = str(metadata.get('universityName', '') or '').strip()
    degree_name = str(metadata.get('degreeName', '') or '').strip()
    word_count = int(metadata.get('wordCount', 0) or 0)

    sorted_sections = sorted(
        [s for s in sections if isinstance(s, dict)],
        key=lambda item: int(item.get('sort_order', 0) or 0)
    )

    total_sections = len(sorted_sections)

    # Cover slide
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_background(slide, COLOR_NAVY)
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(10.45), Inches(0), Inches(2.88), Inches(7.5))
    accent.fill.solid()
    accent.fill.fore_color.rgb = COLOR_ACCENT
    accent.line.fill.background()
    accent2 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(9.55), Inches(5.95), Inches(3.8), Inches(1.55))
    accent2.fill.solid()
    accent2.fill.fore_color.rgb = COLOR_ACCENT_2
    accent2.line.fill.background()
    _add_textbox(slide, Inches(0.82), Inches(0.72), Inches(3.0), Inches(0.28), assignment_type_label.upper(), font_size=11, bold=True, color=COLOR_ACCENT_3)
    _add_textbox(slide, Inches(0.8), Inches(1.25), Inches(9.0), Inches(1.4), title, font_size=34, bold=True, color=COLOR_WHITE, font_name=FONT_HEADLINE)
    subtitle_bits = [assignment_type_label]
    if degree_name:
        subtitle_bits.append(degree_name)
    if word_count:
        subtitle_bits.append(f'{word_count:,} words')
    _add_textbox(slide, Inches(0.84), Inches(2.7), Inches(8.6), Inches(0.45), '  |  '.join(subtitle_bits), font_size=16, color=RGBColor(0xD8, 0xE1, 0xF2))

    if student_name or university_name:
        owner = '  |  '.join([part for part in [student_name, university_name] if part])
        _add_textbox(slide, Inches(0.84), Inches(3.15), Inches(8.4), Inches(0.38), owner, font_size=13, color=RGBColor(0xB7, 0xC6, 0xE1))

    _add_card(slide, Inches(0.82), Inches(4.0), Inches(7.75), Inches(1.35), fill=RGBColor(0x14, 0x26, 0x4A), line=RGBColor(0x2B, 0x43, 0x74), radius=True)
    _add_textbox(
        slide,
        Inches(1.05),
        Inches(4.35),
        Inches(7.15),
        Inches(0.55),
        'This export is structured for slides first, with notes preserved separately.',
        font_size=15,
        color=RGBColor(0xE7, 0xEE, 0xFF),
    )

    # Agenda slide
    if sorted_sections:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        _set_background(slide, COLOR_SOFT)
        _add_top_bar(slide, prs, COLOR_ACCENT)
        _add_textbox(slide, Inches(0.7), Inches(0.55), Inches(6.5), Inches(0.5), 'Agenda', font_size=24, bold=True, color=COLOR_INK, font_name=FONT_HEADLINE)
        agenda = []
        for idx, sec in enumerate(sorted_sections, start=1):
            agenda.append(f'{idx}. {str(sec.get("title", "Section") or "Section").strip()}')
        left_col = agenda[: (len(agenda) + 1) // 2]
        right_col = agenda[(len(agenda) + 1) // 2:]
        _add_card(slide, Inches(0.72), Inches(1.25), Inches(5.95), Inches(5.55), fill=COLOR_WHITE, line=RGBColor(0xD8, 0xE1, 0xF2), radius=True)
        _add_card(slide, Inches(6.76), Inches(1.25), Inches(5.85), Inches(5.55), fill=COLOR_WHITE, line=RGBColor(0xD8, 0xE1, 0xF2), radius=True)
        for col_index, items in enumerate([left_col, right_col]):
            x = 0.98 if col_index == 0 else 7.02
            for idx, item in enumerate(items):
                y = 1.55 + (idx * 0.92)
                bullet = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y + 0.04), Inches(0.22), Inches(0.22))
                bullet.fill.solid()
                bullet.fill.fore_color.rgb = COLOR_ACCENT if col_index == 0 else COLOR_ACCENT_2
                bullet.line.fill.background()
                _add_textbox(slide, Inches(x + 0.32), Inches(y), Inches(5.3), Inches(0.42), item, font_size=18, color=COLOR_INK)
        _add_footer(slide, title, 2, total_sections + 2)

    # Content slides (slide-first rendering with notes preserved separately)
    for sec_index, sec in enumerate(sorted_sections, start=1):
        section_title = str(sec.get('title', '') or 'Section').strip()
        cleaned_title = _strip_slide_prefix(section_title) or section_title

        content_md = str(sec.get('content_markdown', '') or '')
        visible_content, speaker_notes = _extract_speaker_notes(content_md)

        # Determine if this section explicitly indicates a slide (prefer single slide)
        is_explicit_slide = bool(re.search(r'(?i)\bslide\b|\btitle slide\b', section_title))

        blocks = _markdown_to_slide_blocks(visible_content, slide_title=cleaned_title)

        # Group paragraph-like blocks to avoid tiny-sentence splitting
        grouped_lines = []
        buf = []
        for kind, text in blocks:
            if kind in ('bullet', 'number', 'label'):
                if buf:
                    grouped_lines.append(' '.join(buf).strip())
                    buf = []
                if kind == 'bullet':
                    grouped_lines.append(f'• {text}')
                else:
                    grouped_lines.append(text)
            else:
                buf.append(text)
        if buf:
            grouped_lines.append(' '.join(buf).strip())

        # If explicit slide or reasonably short, keep a single slide for the section
        if is_explicit_slide or (len(grouped_lines) <= 6 and sum(len(l) for l in grouped_lines) <= 1200):
            chunks = [grouped_lines]
        else:
            # Use higher thresholds to avoid producing too many slides
            chunks = _chunk_lines(grouped_lines, max_lines=12, max_chars=1400)

        for chunk_index, chunk in enumerate(chunks, start=1):
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            _set_background(slide, COLOR_SOFT)
            _add_top_bar(slide, prs, COLOR_ACCENT if sec_index % 2 else COLOR_ACCENT_2)
            heading = cleaned_title if len(chunks) == 1 else f'{cleaned_title} - Part {chunk_index}'
            _add_textbox(slide, Inches(0.72), Inches(0.56), Inches(10.4), Inches(0.5), heading, font_size=24, bold=True, color=COLOR_INK, font_name=FONT_HEADLINE)
            _add_section_tag(slide, f'Section {sec_index}', color=COLOR_ACCENT_3 if sec_index % 2 else COLOR_ACCENT_2)

            _add_card(slide, Inches(0.72), Inches(1.2), Inches(11.85), Inches(5.45), fill=COLOR_WHITE, line=RGBColor(0xD8, 0xE1, 0xF2), radius=True)
            content_box = slide.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(11.0), Inches(4.95))
            tf = content_box.text_frame
            tf.word_wrap = True
            tf.clear()
            if not chunk:
                p = tf.paragraphs[0]
                run = p.add_run()
                run.text = 'Add content to this section.'
                run.font.name = FONT_BODY
                run.font.size = Pt(18)
                run.font.color.rgb = COLOR_MUTED
            else:
                for idx, line in enumerate(chunk):
                    p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
                    p.space_before = Pt(2)
                    p.space_after = Pt(2)
                    run = p.add_run()
                    if line.startswith('• '):
                        p.level = 0
                        run.text = line[2:]
                        run.font.size = Pt(18 if len(line) < 120 else 16)
                    elif re.match(r'^\d+\.\s+', line):
                        run.text = line
                        run.font.size = Pt(18 if len(line) < 120 else 16)
                    elif line.endswith(':') and len(line) <= 120:
                        run.text = line
                        run.font.bold = True
                        run.font.size = Pt(18)
                    else:
                        run.text = line
                        run.font.size = Pt(20 if len(line) < 120 else 16)
                    run.font.name = FONT_BODY
                    run.font.color.rgb = COLOR_INK

            # Add speaker notes (if present)
            if speaker_notes:
                try:
                    notes_slide = slide.notes_slide
                    notes_tf = notes_slide.notes_text_frame
                    notes_tf.clear()
                    notes_tf.text = _strip_inline_markdown(speaker_notes)
                except Exception:
                    pass

            _add_footer(slide, title, len(prs.slides), total_sections + 2)

    buffer = io.BytesIO()
    prs.save(buffer)
    buffer.seek(0)
    return buffer