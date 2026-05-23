"""
utils/docx_generator.py

Convert assignment sections (markdown) to formatted DOCX.
Uses python-docx (python-docx>=1.1.0 must be in requirements.txt).

Design principles:
- Times New Roman 12pt body (academic standard)
- Heading 1 for document title, Heading 2 for section headings
- 1 inch margins all around
- Page numbers at bottom center
- Student name + university in header
- "AI-Assisted Draft — Review Before Submission" disclaimer on first page
"""
import io
import re
from docx import Document as DocxDocument
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def _add_page_number(paragraph):
    """Add a page number field to a paragraph."""
    run = paragraph.add_run()
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = 'PAGE'
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)


def _markdown_to_paragraphs(doc, markdown_text):
    """
    Convert basic markdown text to DOCX paragraphs.
    Handles: ## headings, **bold**, *italic*, bullet lists, numbered lists, plain text.
    Tables are simplified to text (full table support is Phase 2).
    """
    if not markdown_text:
        return

    lines = markdown_text.strip().split('\n')
    in_list = False

    for line in lines:
        line = line.rstrip()

        # Skip section ## headings (already added as Heading 2 by caller)
        if line.startswith('## '):
            continue

        # ### Sub-headings
        if line.startswith('### '):
            p = doc.add_paragraph(style='Heading 3')
            p.add_run(line[4:].strip())
            in_list = False
            continue

        # Horizontal rule ---
        if re.match(r'^[-─]{3,}$', line):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            in_list = False
            continue

        # Markdown table row (basic — render as normal text in Phase 1)
        if line.startswith('|'):
            # Strip pipe chars and render as tab-separated line
            cells = [c.strip() for c in line.strip('|').split('|')]
            p = doc.add_paragraph('\t'.join(cells))
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            in_list = False
            continue

        # Bullet list
        bullet_match = re.match(r'^[-*•]\s+(.+)', line)
        if bullet_match:
            p = doc.add_paragraph(style='List Bullet')
            _add_inline_formatting(p, bullet_match.group(1))
            in_list = True
            continue

        # Numbered list
        num_match = re.match(r'^\d+\.\s+(.+)', line)
        if num_match:
            p = doc.add_paragraph(style='List Number')
            _add_inline_formatting(p, num_match.group(1))
            in_list = True
            continue

        # Empty line
        if not line.strip():
            if in_list:
                in_list = False
            else:
                doc.add_paragraph('')
            continue

        # Normal paragraph
        in_list = False
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.space_before = Pt(0)
        _add_inline_formatting(p, line)


def _add_inline_formatting(paragraph, text):
    """
    Parse inline **bold** and *italic* and add formatted runs.
    """
    # Pattern: **bold**, *italic*, plain text
    parts = re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith('*') and part.endswith('*'):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            paragraph.add_run(part)


def generate_docx_from_sections(sections, metadata):
    """
    Main entry point. Returns a BytesIO object containing the DOCX file.

    Args:
        sections: list of dicts with keys: section_key, title, content_markdown, sort_order
        metadata: dict with keys: title, assignmentType, studentName, universityName, degreeName, wordCount

    Returns:
        io.BytesIO — ready to send_file
    """
    doc = DocxDocument()

    # ── Page setup: 1 inch margins ───────────────────────────────────────
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # ── Default font: Times New Roman 12pt ───────────────────────────────
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)

    # ── Header: Student Name + University ────────────────────────────────
    header = doc.sections[0].header
    header_para = header.paragraphs[0]
    header_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    student_name = str(metadata.get('studentName', '') or '')
    university = str(metadata.get('universityName', '') or '')
    header_text = f'{student_name}  |  {university}' if student_name or university else 'Assignment Studio — ReportGen'
    header_para.add_run(header_text).font.size = Pt(9)

    # ── Footer: Page number ───────────────────────────────────────────────
    footer = doc.sections[0].footer
    footer_para = footer.paragraphs[0]
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_page_number(footer_para)

    # ── Title ────────────────────────────────────────────────────────────
    title_para = doc.add_paragraph(style='Heading 1')
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_para.add_run(str(metadata.get('title', 'Assignment') or 'Assignment'))
    title_run.font.size = Pt(18)

    # ── Metadata line ─────────────────────────────────────────────────────
    assignment_type_label = {
        'essay': 'Essay',
        'research_paper': 'Research Paper',
        'lab_report': 'Lab Report',
        'case_study': 'Case Study',
        'presentation_script': 'Presentation Script',
    }.get(str(metadata.get('assignmentType', '') or ''), 'Assignment')

    meta_para = doc.add_paragraph()
    meta_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_text = f'{assignment_type_label}'
    if metadata.get('degreeName'):
        meta_text += f'  |  {metadata["degreeName"]}'
    if metadata.get('wordCount'):
        meta_text += f'  |  {metadata["wordCount"]:,} words'
    meta_run = meta_para.add_run(meta_text)
    meta_run.font.size = Pt(10)
    meta_run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    # ── Disclaimer ────────────────────────────────────────────────────────
    disclaimer_para = doc.add_paragraph()
    disclaimer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    d_run = disclaimer_para.add_run('⚠ AI-Assisted Draft — Review Before Submission')
    d_run.font.size = Pt(9)
    d_run.font.italic = True
    d_run.font.color.rgb = RGBColor(0x94, 0xa3, 0xb8)
    disclaimer_para.paragraph_format.space_after = Pt(12)

    # ── Section divider ──────────────────────────────────────────────────
    doc.add_paragraph()

    # ── Sections ─────────────────────────────────────────────────────────
    sorted_sections = sorted(
        [s for s in sections if isinstance(s, dict)],
        key=lambda x: int(x.get('sort_order', 0) or 0)
    )

    for sec in sorted_sections:
        title = str(sec.get('title', '') or '')
        content = str(sec.get('content_markdown', '') or '')

        if title:
            heading = doc.add_paragraph(style='Heading 2')
            heading.add_run(title)

        if content:
            _markdown_to_paragraphs(doc, content)

        # Space between sections
        doc.add_paragraph()

    # ── Save to BytesIO ──────────────────────────────────────────────────
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
